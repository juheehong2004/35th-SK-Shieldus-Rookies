"""
[응용-2] 보안 뉴스 자동 크롤링 및 엑셀 리포트 메일링 파이프라인
- requests + BeautifulSoup : 보안 뉴스 기사 제목/링크 수집
- openpyxl                 : 위험도별 조건부 서식이 적용된 엑셀 리포트 저장
- smtplib + email.mime     : 엑셀 첨부 + 고위험 기사 HTML 표 메일 발송
- python-dotenv            : 계정 정보를 .env 로 분리 관리

[.env 예시]
GOOGLE_EMAIL_USER=google_id@gmail.com
GOOGLE_EMAIL_PASS=앱비밀번호16자리
RECEIVER_EMAIL=teamleader@company.com
"""

import logging
import os
import smtplib
from datetime import datetime
from email.mime.application import MIMEApplication
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from html import escape
from urllib.parse import urljoin, urlparse

import requests
from bs4 import BeautifulSoup
from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

# 로그 출력 형식 설정 (시간 + 로그 레벨 + 메시지)
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
)

# ---------------------------------------------------------------------------
# 설정값
# ---------------------------------------------------------------------------
TARGET_URL = "https://www.boannews.com/"
TARGET_DOMAIN = "boannews.com"
# 브라우저 요청처럼 보이도록 User-Agent 추가
HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
        "(KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
    )
}
MIN_TITLE_LENGTH = 15 # 제목 15자 미만은 메뉴/배너로 보고 제외
RISK_KEYWORDS = ["취약점", "유출", "해킹"] # 위험도 High 판단 키워드


# ---------------------------------------------------------------------------
# 1. 데이터 수집 엔진 (requests + BeautifulSoup)
# ---------------------------------------------------------------------------
def judge_risk(title):
    """제목에 위험 키워드가 하나라도 있으면 High, 아니면 Normal"""
    for keyword in RISK_KEYWORDS:
        if keyword in title:
            return "High"
    return "Normal"


def collect_articles():
    """보안 뉴스 메인 페이지에서 기사 제목/링크/위험도를 수집한다."""
    # 페이지 요청 (네트워크 오류 발생 시 로그만 남기고 빈 리스트 반환)
    try:
        response = requests.get(TARGET_URL, headers=HEADERS, timeout=30)
        response.raise_for_status()
    except requests.exceptions.RequestException as e:
        logging.error(f"[수집 오류] 페이지 요청 실패: {e}")
        return []

    # 한글 깨짐 방지를 위해 인코딩 자동 감지 후 HTML 파싱
    response.encoding = response.apparent_encoding
    soup = BeautifulSoup(response.text, "html.parser")

    articles = []

    # 중복 기사 판별용
    seen_titles = set()
    seen_links = set()

    # class 이름에 의존하지 않고 모든 링크를 훑은 뒤 조건으로 필터링
    for a in soup.find_all("a"):
        href = a.get("href")
        title = a.get_text(strip=True)

        # 링크가 없거나 의미 없는 링크는 제외
        if not href or href.startswith(("#", "javascript:")):
            continue

        # 제목 15자 미만 제외 (메뉴/배너)
        if len(title) < MIN_TITLE_LENGTH:
            continue

        # urljoin 으로 절대 링크 생성 후 대상 사이트 도메인만 채택
        link = urljoin(TARGET_URL, href)
        if TARGET_DOMAIN not in urlparse(link).netloc:
            continue

        # 중복 기사 제외
        if title in seen_titles or link in seen_links:
            continue
        seen_titles.add(title)
        seen_links.add(link)

        # 기사 정보 저장 (위험도는 제목 기준으로 판별)
        articles.append(
            {
                "no": len(articles) + 1,
                "title": title,
                "link": link,
                "risk": judge_risk(title),
            }
        )

    logging.info(f"기사 {len(articles)}건 수집 완료")
    return articles


# ---------------------------------------------------------------------------
# 2. 엑셀 리포트 생성 (openpyxl)
# ---------------------------------------------------------------------------
def save_excel(articles):
    """수집 결과를 엑셀로 저장하고 파일명을 반환한다."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Security News"

    # 헤더 행 추가 후 배경색/글자색으로 강조
    ws.append(["순번", "제목", "링크", "위험도"])
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font

    # 기사 데이터 행 추가 (위험도 High 이면 빨간색 + 굵게)
    for article in articles:
        ws.append([article["no"], article["title"], article["link"], article["risk"]])
        if article["risk"] == "High":
            for cell in ws[ws.max_row]:
                cell.font = Font(color="FF0000", bold=True)

    # 파일명에 당일 날짜 포함: Security_Report_YYYYMMDD.xlsx
    filename = f"Security_Report_{datetime.now().strftime('%Y%m%d')}.xlsx"
    wb.save(filename)
    logging.info(f"엑셀 저장 완료: {filename}")
    return filename


# ---------------------------------------------------------------------------
# 3. 이메일 발송 (smtplib + MIMEMultipart + 첨부파일)
# ---------------------------------------------------------------------------
def build_html_body(articles):
    """위험도 High 기사만 모아 클릭 가능한 HTML 표로 구성한다."""
    high_articles = [a for a in articles if a["risk"] == "High"]

    if not high_articles:
        return "<p>오늘 위험도 High 기사는 없습니다.</p>"

    # 기사별 표 행(<tr>) 생성 (제목은 <a> 태그로 클릭 가능하게 처리)
    rows = ""
    for a in high_articles:
        rows += (
            "<tr>"
            f"<td style='border:1px solid #ccc; padding:6px;'>{a['no']}</td>"
            f"<td style='border:1px solid #ccc; padding:6px;'>"
            f"<a href='{escape(a['link'])}'>{escape(a['title'])}</a></td>"
            f"<td style='border:1px solid #ccc; padding:6px; color:red;'><b>{a['risk']}</b></td>"
            "</tr>"
        )

    # 최종 HTML 본문 구성
    html = f"""
    <html>
      <body>
        <p>안녕하세요. 오늘 수집된 보안 뉴스 중 <b>위험도 High</b> 기사는 {len(high_articles)}건입니다.</p>
        <table style="border-collapse:collapse;">
          <tr style="background-color:#4472C4; color:white;">
            <th style="border:1px solid #ccc; padding:6px;">순번</th>
            <th style="border:1px solid #ccc; padding:6px;">제목</th>
            <th style="border:1px solid #ccc; padding:6px;">위험도</th>
          </tr>
          {rows}
        </table>
        <p>전체 기사 목록은 첨부된 엑셀 파일을 확인해주세요.</p>
      </body>
    </html>
    """
    return html


def send_email(articles, excel_filename):
    """엑셀 파일을 첨부한 HTML 메일을 발송한다."""
    # .env 에서 계정 정보 로드
    email_id = os.getenv("GOOGLE_EMAIL_USER")
    email_pw = os.getenv("GOOGLE_EMAIL_PASS")
    receiver = os.getenv("RECEIVER_EMAIL", email_id)

    if not email_id or not email_pw:
        logging.error("[발송 오류] .env 에 GOOGLE_EMAIL_USER / GOOGLE_EMAIL_PASS 가 설정되어 있지 않습니다.")
        return False

    # 메일 기본 정보 설정
    msg = MIMEMultipart()
    msg["From"] = email_id
    msg["To"] = receiver
    msg["Subject"] = f"[보안 뉴스] {datetime.now().strftime('%Y-%m-%d')} 위험도 High 기사 리포트"

    # 본문 구성: HTML 표
    msg.attach(MIMEText(build_html_body(articles), "html"))

    # 파일 첨부: 'rb' 모드로 열어 MIMEApplication 생성 + Content-Disposition 헤더 추가
    try:
        with open(excel_filename, "rb") as f:
            part = MIMEApplication(f.read(), Name=excel_filename)
        part.add_header("Content-Disposition", "attachment", filename=excel_filename)
        msg.attach(part)
    except OSError as e:
        logging.error(f"[첨부 오류] 엑셀 파일을 열 수 없습니다: {e}")
        return False

    # 발송 처리: smtp.gmail.com(587) 연결 -> starttls() 암호화 -> 로그인 -> send_message() 발송
    try:
        with smtplib.SMTP("smtp.gmail.com", 587, timeout=30) as server:
            server.starttls()
            server.login(email_id, email_pw)
            server.send_message(msg)
        logging.info(f"메일 발송 완료 -> {receiver}")
        return True
    except smtplib.SMTPAuthenticationError as e:
        logging.error(f"[발송 오류] 로그인 실패 (ID/앱 비밀번호 확인): {e}")
    except (smtplib.SMTPException, OSError) as e:
        logging.error(f"[발송 오류] 메일 전송 실패: {e}")
    return False


# ---------------------------------------------------------------------------
# 4. 통합 파이프라인
# ---------------------------------------------------------------------------
def main():
    load_dotenv() # .env 파일에서 계정 정보 로드

    # 1) 기사 수집 (수집 결과가 없으면 비정상 종료 없이 종료)
    articles = collect_articles()
    if not articles:
        logging.warning("수집된 기사가 없어 파이프라인을 종료합니다.")
        return

    # 2) 엑셀 리포트 저장
    try:
        excel_filename = save_excel(articles)
    except OSError as e:
        logging.error(f"[저장 오류] 엑셀 파일 저장 실패: {e}")
        return

    # 3) 메일 발송
    send_email(articles, excel_filename)


if __name__ == "__main__":
    main()